from __future__ import annotations

import argparse
import csv
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any, Iterable, Mapping

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tools.azure_sql_migration.session_time_corrections import FORMAT, canonical_time
from tools.azure_sql_migration.session_time_reconciliation import normalize_session_type, normalize_topic, reconcile_sessions


def sha256_file(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_current_csv(path: str | Path) -> list[dict[str, Any]]:
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        source = list(csv.DictReader(handle))
    rows = []
    for row in source:
        rows.append({
            "SessionId": str(row.get("ID") or "").strip().upper(),
            "Date": str(row.get("Date") or "").strip()[:10],
            "CourseCode": str(row.get("Course") or "").strip(),
            "SessionType": str(row.get("Type") or "").strip(),
            "PAWS_Topic": str(row.get("Topic") or "").strip(),
            "CurrentStart": canonical_time(row.get("Start"), field="current Start"),
            "CurrentEnd": canonical_time(row.get("End"), field="current End"),
        })
    if len({row["SessionId"] for row in rows}) != len(rows):
        raise ValueError("current session CSV contains duplicate SessionId values")
    return rows


def load_timetable_xlsx(path: str | Path) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the authoritative timetable") from exc
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "Timetable" not in workbook.sheetnames:
        raise ValueError("authoritative timetable has no Timetable sheet")
    sheet = workbook["Timetable"]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(iterator)]
    rows = []
    for values in iterator:
        source = dict(zip(headers, values))
        date = str(source.get("Date") or "").strip()[:10]
        if not date.startswith("20"):
            continue
        if not all(str(source.get(name) or "").strip() for name in ("Course", "Type", "Topic", "Start Time", "End Time")):
            continue
        rows.append({
            "Date": date,
            "CourseCode": str(source["Course"]).strip(),
            "SessionType": str(source["Type"]).strip(),
            "TimetableTopic": str(source["Topic"]).strip(),
            "TimetableStart": canonical_time(source["Start Time"], field="timetable Start Time"),
            "TimetableEnd": canonical_time(source["End Time"], field="timetable End Time"),
        })
    return rows


def load_audit_xlsx(path: str | Path) -> dict[tuple[str, str, str, str], list[dict[str, Any]]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the supporting audit") from exc
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if "All Comparison" not in workbook.sheetnames:
        raise ValueError("supporting audit has no All Comparison sheet")
    sheet = workbook["All Comparison"]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(iterator)]
    index: dict[tuple[str, str, str, str], list[dict[str, Any]]] = {}
    for values in iterator:
        row = dict(zip(headers, values))
        date_value = row.get("Date Timetable")
        date = date_value.strftime("%Y-%m-%d") if hasattr(date_value, "strftime") else str(date_value or "")[:10]
        key = (
            date,
            str(row.get("Course") or "").strip(),
            normalize_session_type(row.get("Type Timetable")),
            normalize_topic(row.get("Topic Timetable")),
        )
        if all(key):
            index.setdefault(key, []).append(row)
    return index


def add_audit_evidence(rows: list[dict[str, Any]], audit_index: Mapping[tuple[str, str, str, str], list[dict[str, Any]]]) -> None:
    for row in rows:
        key = (
            str(row.get("Date") or "")[:10],
            str(row.get("CourseCode") or "").strip(),
            normalize_session_type(row.get("SessionType")),
            normalize_topic(row.get("TimetableTopic") or row.get("PAWS_Topic")),
        )
        candidates = audit_index.get(key, [])
        exact = [candidate for candidate in candidates if str(candidate.get("Match Method") or "").casefold().startswith("exact")]
        row["AuditMatched"] = "YES" if exact else "NO"
        row["AuditMatchMethod"] = str(exact[0].get("Match Method") or "") if exact else ""
        row["AuditConfidence"] = str(exact[0].get("Confidence") or "") if exact else ""
        if exact:
            row["Evidence"] = f"{row['Evidence']}; supporting audit: {row['AuditMatchMethod']} (confidence {row['AuditConfidence']})"


def build_manifest(rows: Iterable[Mapping[str, Any]], sources: Mapping[str, Any]) -> dict[str, Any]:
    corrections = []
    for row in rows:
        if row.get("CorrectionRequired") != "YES":
            continue
        corrections.append({
            "sessionId": str(row["SessionId"]).lower(),
            "oldStart": row["CurrentStart"],
            "oldEnd": row["CurrentEnd"],
            "newStart": row["ProposedStart"],
            "newEnd": row["ProposedEnd"],
            "source": sources["master"]["fileName"],
            "sourceDate": row["Date"],
            "sourceCourse": row["CourseCode"],
            "sourceTopic": row["TimetableTopic"],
            "sourceMatchMethod": row["MatchMethod"],
        })
    corrections.sort(key=lambda item: item["sessionId"])
    return {"format": FORMAT, "generatedFrom": dict(sources), "corrections": corrections}


def _seconds(value: str) -> int:
    hour, minute, second = (int(part) for part in value.split(":"))
    return hour * 3600 + minute * 60 + second


def summarize_reconciliation(rows: Iterable[Mapping[str, Any]]) -> dict[str, Any]:
    materialized = list(rows)
    classifications = Counter(str(row["Classification"]) for row in materialized)
    sentinel = lambda row: row["CurrentStart"] == "00:00:00" and row["CurrentEnd"] == "00:00:00"
    reversed_interval = lambda row: not sentinel(row) and _seconds(row["CurrentStart"]) >= _seconds(row["CurrentEnd"])
    positionable = lambda row: _seconds(row["CurrentStart"]) >= 27000 and _seconds(row["CurrentEnd"]) <= 61200 and _seconds(row["CurrentStart"]) < _seconds(row["CurrentEnd"])
    early = lambda row: not sentinel(row) and not reversed_interval(row) and _seconds(row["CurrentStart"]) < 27000
    deterministic = classifications["DETERMINISTIC_TIME_CORRECTION"] + classifications["UNKNOWN_TIME_CONFIRMED_BY_TIMETABLE"]
    return {
        "totalSessions": len(materialized),
        "uniqueSessionIds": len({str(row["SessionId"]).lower() for row in materialized if row.get("SessionId")}),
        "currentGridPositionable": sum(positionable(row) for row in materialized),
        "currentEarlyClock": sum(early(row) for row in materialized),
        "currentReversed": sum(reversed_interval(row) for row in materialized),
        "currentMidnightSentinel": sum(sentinel(row) for row in materialized),
        "deterministicSourceMatches": sum(row.get("TimetableMatched") == "YES" for row in materialized),
        "deterministicCorrections": deterministic,
        "alreadyCorrect": classifications["MATCH_ALREADY_CORRECT"],
        "ambiguous": classifications["SOURCE_AMBIGUOUS"] + classifications["MULTIPLE_TIMETABLE_MATCHES"],
        "unmatched": classifications["NO_TIMETABLE_MATCH"],
        "legitimateOutOfWindow": classifications["LEGITIMATE_OUT_OF_WINDOW"],
        "otherDataDefect": classifications["OTHER_DATA_DEFECT"],
        "classifications": dict(sorted(classifications.items())),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Reconcile PAWS session times to an authoritative UCVM timetable.")
    parser.add_argument("--current-csv", required=True)
    parser.add_argument("--master-xlsx", required=True)
    parser.add_argument("--audit-xlsx")
    parser.add_argument("--reconciliation-json", required=True)
    parser.add_argument("--manifest-output", required=True)
    args = parser.parse_args()

    current_path = Path(args.current_csv).resolve()
    master_path = Path(args.master_xlsx).resolve()
    audit_path = Path(args.audit_xlsx).resolve() if args.audit_xlsx else None
    current = load_current_csv(current_path)
    timetable = load_timetable_xlsx(master_path)
    reconciliation = reconcile_sessions(current, timetable)
    sources: dict[str, Any] = {
        "master": {"fileName": master_path.name, "sha256": sha256_file(master_path)},
        "currentSessions": {"fileName": current_path.name, "sha256": sha256_file(current_path)},
    }
    if audit_path:
        sources["supportingAudit"] = {"fileName": audit_path.name, "sha256": sha256_file(audit_path)}
    manifest = build_manifest(reconciliation, sources)
    if audit_path:
        add_audit_evidence(reconciliation, load_audit_xlsx(audit_path))
    summary = summarize_reconciliation(reconciliation)
    summary["supportingAuditMatches"] = sum(row.get("AuditMatched") == "YES" for row in reconciliation)
    payload = {"sources": sources, "summary": summary, "sessions": reconciliation}

    reconciliation_path = Path(args.reconciliation_json).resolve()
    manifest_path = Path(args.manifest_output).resolve()
    reconciliation_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    reconciliation_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"reconciliation": str(reconciliation_path), "manifest": str(manifest_path), "summary": payload["summary"]}, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
