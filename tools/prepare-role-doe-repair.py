"""Read-only source reconciliation for legacy role DOE recovery.

Inputs and output contain confidential Faculty data and must stay outside Git.
No SQL connection, source workbook recalculation, or live write is performed.
"""
import argparse
from collections import defaultdict, Counter
from datetime import date, datetime
import hashlib
import json
import math
from pathlib import Path
import re
import unicodedata
import uuid
from openpyxl import load_workbook

LEGACY = 'Course Coordinator / HICC (legacy ambiguous)'

def key(value):
    return ' '.join(unicodedata.normalize('NFC', str(value or '')).split()).casefold()

def course_key(value):
    # The legacy import used VTMD for the workbook's "Vet Med" course labels.
    return re.sub(r'^vet med\s+', 'vtmd ', key(value))

def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)

def iso(value):
    return value.date().isoformat() if isinstance(value, datetime) else value.isoformat() if isinstance(value, date) else None

def prepare(workbook_path, snapshot):
    digest = hashlib.sha256(Path(workbook_path).read_bytes()).hexdigest()
    names = defaultdict(set)
    faculty = {r['FacultyId'].upper(): r for r in snapshot['faculty']}
    for fid, row in faculty.items():
        names[key(f"{row.get('PreferredLastName')}, {row.get('PreferredFirstName')}")].add(fid)
        names[key(row.get('DisplayName'))].add(fid)
    for row in snapshot['aliases']:
        names[key(row['AliasName'])].add(row['FacultyId'].upper())
    # Existing resolved source names supply an audited crosswalk, including HR
    # names that differ from the displayed preferred name. Reject collisions.
    for row in snapshot['roles']:
        if row['FacultyId'] and row.get('SourceFacultyName'):
            names[key(row['SourceFacultyName'])].add(row['FacultyId'].upper())
    values = load_workbook(workbook_path, read_only=True, data_only=True)
    formulas = load_workbook(workbook_path, read_only=True, data_only=False)
    evidence, rejected = [], []
    for sheet, role_type, name_col, doe_col, start_col, end_col, max_row in [
        ('Course Coordinator', 'Course Coordinator', 2, 4, 8, 9, 116),
        ('HICCs', 'HICC', 2, 5, 9, 10, 47),
        ('CCC', 'CCC', 14, 19, None, None, 40),
    ]:
        ws, fs = values[sheet], formulas[sheet]
        cached_rows = ws.iter_rows(min_row=3 if sheet != 'CCC' else 2, max_row=max_row, values_only=True)
        formula_rows = fs.iter_rows(min_row=3 if sheet != 'CCC' else 2, max_row=max_row, values_only=True)
        for row_num, (row, formula_row) in enumerate(zip(cached_rows, formula_rows), 3 if sheet != 'CCC' else 2):
            name = row[name_col-1]
            if not isinstance(name, str) or not name.strip():
                continue
            assigned = 'CCC shifts' if sheet == 'CCC' else row[0]
            if not assigned:
                continue
            ids = names.get(key(name), set())
            result = row[doe_col-1]
            # A genuinely blank CC DOE cell is not an explicit zero. In these
            # rows the separately populated HICC row is the credit-bearing role.
            if not numeric(result):
                rejected.append({'sheet': sheet, 'row': row_num, 'name': name, 'assignment': assigned, 'reason': 'source-doe-not-numeric', 'value': result})
                continue
            if len(ids) != 1 or next(iter(ids)) not in faculty:
                rejected.append({'sheet': sheet, 'row': row_num, 'name': name, 'assignment': assigned, 'reason': 'identity-not-unique', 'candidates': sorted(ids)})
                continue
            fid = next(iter(ids))
            if not faculty[fid]['Active']:
                rejected.append({'sheet': sheet, 'row': row_num, 'name': name, 'assignment': assigned, 'reason': 'inactive-faculty'})
                continue
            evidence.append({
                'facultyId': fid, 'sourceFacultyName': name, 'roleType': role_type,
                'assignment': str(assigned).strip(), 'doe': float(result),
                'sheet': sheet, 'row': row_num, 'cell': f'{"S" if sheet == "CCC" else "D" if sheet == "Course Coordinator" else "E"}{row_num}',
                'formula': formula_row[doe_col-1] if isinstance(formula_row[doe_col-1], str) and formula_row[doe_col-1].startswith('=') else None,
                'effectiveDate': iso(row[start_col-1]) if start_col else None,
                'expirationDate': iso(row[end_col-1]) if end_col else None,
                'sourceReference': f'{Path(workbook_path).name} sha256:{digest} {sheet}!row{row_num}',
                'sourceNotes': ' | '.join(str(v) for v in ([row[4],row[11]] if sheet == 'Course Coordinator' else [row[5],row[13]] if sheet == 'HICCs' else [f"shifts={row[16]}",f"halfDays={row[19]}",f"halfDayDeduction={row[20]}"]) if v is not None),
            })
    by_business = defaultdict(list)
    for row in evidence:
        by_business[(row['facultyId'], course_key(row['assignment']), row['roleType'])].append(row)
    duplicates = [list(k) for k, rows in by_business.items() if len(rows) > 1]
    if duplicates:
        raise ValueError(f'Duplicate role source keys require explicit reconciliation: {duplicates}')
    operations, unresolved, consumed = [], [], set()
    for before in snapshot['roles']:
        if before['RoleType'] not in (LEGACY, 'CCC'):
            continue
        fid = str(before['FacultyId']).upper()
        assigned = course_key(before['CourseOrSubjectOrRotation'])
        kinds = ['Course Coordinator', 'HICC'] if before['RoleType'] == LEGACY else ['CCC']
        candidates = [rows[0] for kind in kinds if (rows := by_business.get((fid, assigned, kind)))]
        if not candidates:
            unresolved.append({'roleAssignmentId': before['RoleAssignmentId'], 'reason': 'no-numeric-source-match', 'facultyId': fid, 'assignment': before['CourseOrSubjectOrRotation']})
            continue
        if before['DoeMappedValue'] is not None or before['DoeOverrideValue'] is not None:
            unresolved.append({'roleAssignmentId': before['RoleAssignmentId'], 'reason': 'existing-value-preserved'})
            continue
        for index, source in enumerate(candidates):
            source_key = (source['facultyId'], course_key(source['assignment']), source['roleType'])
            consumed.add(source_key)
            # Stable identity for the additional separately credited HICC row.
            new_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"https://paws.ucalgary.ca/role-doe-repair/{digest}/{source['sheet']}/{source['row']}"))
            operations.append({'action': 'update' if index == 0 else 'insert', 'roleAssignmentId': before['RoleAssignmentId'] if index == 0 else new_id, 'before': before, 'source': source})
    unmatched = [r for r in evidence if (r['facultyId'],course_key(r['assignment']),r['roleType']) not in consumed]
    remaining = []
    for source in unmatched:
        # The legacy import omitted CCC participants with an explicitly recorded
        # zero. Preserve them as real zero-valued lines, not absent assignments.
        existing = [r for r in snapshot['roles'] if str(r['FacultyId']).upper() == source['facultyId']
                    and r['RoleType'] == source['roleType']
                    and course_key(r['CourseOrSubjectOrRotation']) == course_key(source['assignment'])]
        if existing:
            remaining.append(source)
            continue
        if source['roleType'] != 'CCC':
            remaining.append(source)
            continue
        new_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"https://paws.ucalgary.ca/role-doe-repair/{digest}/{source['sheet']}/{source['row']}"))
        operations.append({'action': 'insert', 'roleAssignmentId': new_id, 'before': None, 'source': source})
    unmatched = remaining
    return {'academicYear': '2026-27', 'sourceSha256': digest, 'sourceFile': Path(workbook_path).name,
            'operations': operations, 'unresolved': unresolved, 'unmatchedSourceRoles': unmatched, 'rejectedSourceRows': rejected,
            'summary': {'beforeRoleRows': len(snapshot['roles']), 'beforeMissing': sum(r['DoeMappedValue'] is None and r['DoeOverrideValue'] is None for r in snapshot['roles']),
                        'sourceNumericRows': len(evidence), 'operations': dict(Counter(r['action'] for r in operations)),
                        'explicitZeroRows': sum(r['source']['doe'] == 0 for r in operations), 'unresolved': len(unresolved), 'unmatchedSourceRoles': len(unmatched), 'rejectedSourceRows': len(rejected)}}

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', required=True)
    parser.add_argument('--snapshot', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    result = prepare(args.workbook, json.loads(Path(args.snapshot).read_text(encoding='utf-8-sig')))
    Path(args.output).write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(json.dumps(result['summary']))
