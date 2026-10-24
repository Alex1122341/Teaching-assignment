"""Read source workbooks; write a private, deterministic supervision recovery plan.

No workbook edits or SQL writes. Cached DOE must match the source formulas.
The five missing faculty zeros and five missing formula rows were confirmed by
the owner on 2026-10-08. No fuzzy name matching or inferred historical counts.
"""
import argparse
from collections import defaultdict
import hashlib
import json
import math
from pathlib import Path
import unicodedata
import uuid
from openpyxl import load_workbook

def key(v):
    return ' '.join(unicodedata.normalize('NFC', str(v or '')).split()).casefold()

def prepare(master, faculty_master, guideline, snapshot):
    digest = hashlib.sha256(Path(master).read_bytes()).hexdigest()
    names = defaultdict(set)
    faculties = {r['FacultyId']: r for r in snapshot['faculty']}
    for fid, r in faculties.items():
        for n in (r['DisplayName'], f"{r.get('PreferredLastName')}, {r.get('PreferredFirstName')}"):
            names[key(n)].add(fid)
    for r in snapshot['aliases']: names[key(r['AliasName'])].add(r['FacultyId'])
    for r in snapshot['roles']:
        if r.get('FacultyId') and r.get('SourceFacultyName'): names[key(r['SourceFacultyName'])].add(r['FacultyId'])
    hr = load_workbook(faculty_master, data_only=True, read_only=True)
    for r in hr['Key - UCID.Name'].iter_rows(min_row=2, values_only=True):
        ids = names[key(r[0])]
        if len(ids) == 1: names[key(r[1])].update(ids)
    book = load_workbook(master, data_only=True, read_only=True)
    sheet = book['# Trainees']
    assert sheet['A1'].value == 'HR FULL Name' and sheet['AA1'].value == 'Trainee DOE (Sum)', 'Source layout changed'
    rows, exclusions = [], []
    for i, r in enumerate(sheet.iter_rows(values_only=True), 1):
        if i < 3 or not isinstance(r[0], str): continue
        ids = names[key(r[0])]
        if len(ids) != 1:
            exclusions.append({'row': i, 'sourceName': r[0], 'reason': 'no-unique-canonical-identity'})
            continue
        fid = next(iter(ids))
        if not faculties[fid]['Active']: continue
        values = [r[c-1] if r[c-1] is not None else 0 for c in (4,6,8,11,14,16,18)]
        if any(not isinstance(v, (int, float)) or not math.isfinite(v) or v < 0 for v in values):
            raise ValueError(f'Non-numeric trainee count at row {i}')
        u,v,p,im,g,c,d = values
        capped = lambda n: 0 if n < 1 else 0.5 if n == 1 else 1
        credits = [u, capped(v), p*.75, capped(im), g*3, c*2, d*1.5]
        raw = sum(credits)
        if r[26] is not None and abs(float(r[26]) - raw) > .000001:
            raise ValueError(f'Cached source DOE differs from formulas at row {i}')
        if r[26] is None and i not in (14,16,29,50,79):
            raise ValueError(f'Unapproved missing DOE formula at row {i}')
        evidence = {'sourceRow':i, 'sourceName':r[0], 'counts':values, 'categoryCredits':credits,
                    'cachedRawDoe':r[26], 'rawDoe':raw, 'historyBasis':'current-year; no historical counts supplied',
                    'missingFormulaOwnerConfirmed':r[26] is None, 'sourceSummary':r[2]}
        rows.append({'snapshotId':str(uuid.uuid5(uuid.NAMESPACE_URL,f'paws:{digest}:trainees:2026-27:{fid}')),
                     'facultyId':fid, 'rawDoe':raw, 'evidence':evidence,
                     'sourceReference':f'Teaching Assignments 2026 27 MASTER (8).xlsx · # Trainees!A{i}:AA{i}'})
    for display in ('Derrick, Rancourt','Shannon, Toy','Emma, Vaasjo','Tara, Arndt','Douglas, Morck'):
        ids=names[key(display)]
        if len(ids)!=1: raise ValueError('Owner-confirmed faculty identity is not unique')
        fid=next(iter(ids))
        if any(r['facultyId']==fid for r in rows): raise ValueError('Owner-confirmed missing source row is now present')
        rows.append({'snapshotId':str(uuid.uuid5(uuid.NAMESPACE_URL,f'paws:{digest}:trainees:2026-27:{fid}')),
                     'facultyId':fid,'rawDoe':0,'evidence':{'ownerConfirmedZero':'2026-10-08','historyBasis':'current-year'},
                     'sourceReference':'Owner confirmed zero trainee count · 2026-10-08'})
    if len(rows)!=116 or len({r['facultyId'] for r in rows})!=116: raise ValueError('Unexpected source coverage')
    return {'academicYear':'2026-27','sourceSha256':digest,
            'guidelineSha256':hashlib.sha256(Path(guideline).read_bytes()).hexdigest(),
            'identitySourceSha256':hashlib.sha256(Path(faculty_master).read_bytes()).hexdigest(),
            'workloadSourceReference':'UCVM Workload Guidelines · updated 2025-12-23 · section 6.3, Table 2, p. 4',
            'parameters':{'splitThreshold':35,'splitRatio':.5,'highTeachingTraineeCeiling':15,
                          'teachingFocusedTraineeCeiling':7.5,'rollingAverageYears':3},
            'rows':rows,'exclusions':exclusions,
            'summary':{'rows':len(rows),'rawDoe':sum(r['rawDoe'] for r in rows),'explicitZeros':sum(r['rawDoe']==0 for r in rows)}}

if __name__ == '__main__':
    p=argparse.ArgumentParser()
    for name in ('master','faculty-master','guideline','snapshot','output'): p.add_argument('--'+name,required=True)
    a=p.parse_args()
    plan=prepare(a.master,a.faculty_master,a.guideline,json.loads(Path(a.snapshot).read_text(encoding='utf-8-sig')))
    Path(a.output).write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(plan['summary']))
