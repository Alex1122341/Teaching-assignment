from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook


def serialize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, (str, int, float, bool)):
        return value
    text = getattr(value, "text", None)
    if text is not None:
        return str(text)
    return str(value)


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _formula_text(cell: Any) -> str | None:
    value = cell.value
    if getattr(cell, "data_type", None) == "f":
        text = getattr(value, "text", None)
        if text is not None:
            return str(text)
        return str(value) if value is not None else None
    if isinstance(value, str) and value.startswith("="):
        return value
    text = getattr(value, "text", None)
    if text and str(text).startswith("="):
        return str(text)
    return None


def _header_label(ws_formula: Any, ws_values: Any, column: int) -> str:
    f = ws_formula.cell(row=1, column=column).value
    v = ws_values.cell(row=1, column=column).value
    candidate = v if not _is_blank(v) else f
    if _is_blank(candidate):
        return f"column_{column}"
    return str(candidate).strip()


def iter_nonempty_rows(ws_formula: Any, ws_values: Any) -> Iterator[dict[str, Any]]:
    max_row = max(ws_formula.max_row or 0, ws_values.max_row or 0)
    max_column = max(ws_formula.max_column or 0, ws_values.max_column or 0)
    headers = {column: _header_label(ws_formula, ws_values, column) for column in range(1, max_column + 1)}

    for row_number in range(1, max_row + 1):
        cells: list[dict[str, Any]] = []
        row_has_content = False
        for column in range(1, max_column + 1):
            formula_cell = ws_formula.cell(row=row_number, column=column)
            value_cell = ws_values.cell(row=row_number, column=column)
            formula = _formula_text(formula_cell)
            value = value_cell.value if formula is not None else formula_cell.value
            if formula is not None and _is_blank(value_cell.value):
                value = None
            if not _is_blank(formula) or not _is_blank(value):
                row_has_content = True
                cells.append({
                    "column": column,
                    "coordinate": formula_cell.coordinate,
                    "header": headers[column],
                    "value": serialize_cell(value),
                    "formula": formula,
                })
        if row_has_content:
            yield {"sourceRow": row_number, "cells": cells}


@dataclass(frozen=True)
class WorkbookRecord:
    path: Path
    file_name: str
    file_sha256: str
    file_byte_length: int
    sheets: list[dict[str, Any]]
    rows: list[dict[str, Any]]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def inspect_workbook(path: str | Path) -> WorkbookRecord:
    path = Path(path).resolve()
    formula_book = load_workbook(path, data_only=False, read_only=False)
    values_book = load_workbook(path, data_only=True, read_only=False)
    try:
        sheets: list[dict[str, Any]] = []
        rows: list[dict[str, Any]] = []
        for sheet_index, name in enumerate(formula_book.sheetnames, start=1):
            ws_formula = formula_book[name]
            ws_values = values_book[name]
            sheet_rows = list(iter_nonempty_rows(ws_formula, ws_values))
            sheets.append({
                "sheet": name,
                "sheetIndex": sheet_index,
                "maxRow": max(ws_formula.max_row or 0, ws_values.max_row or 0),
                "maxColumn": max(ws_formula.max_column or 0, ws_values.max_column or 0),
                "nonEmptyRows": len(sheet_rows),
            })
            for row in sheet_rows:
                rows.append({
                    "workbook": path.name,
                    "sheet": name,
                    "sheetIndex": sheet_index,
                    **row,
                })
        return WorkbookRecord(
            path=path,
            file_name=path.name,
            file_sha256=sha256_file(path),
            file_byte_length=path.stat().st_size,
            sheets=sheets,
            rows=rows,
        )
    finally:
        formula_book.close()
        values_book.close()
